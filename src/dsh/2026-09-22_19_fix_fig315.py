# -*- coding: utf-8 -*-
"""补回说明书的 图 3-15 配图（体检与标定摘要）。当前文件里 3-15 与 3-16 共用同一张图，
这里把 3-15 那处换成 shots/s7_report_metro_kpi.png，其余内容一律不动。
用法：python src/dsh/2026-09-22_19_fix_fig315.py
"""
import io, re, zipfile
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX = 'deliverables/澜脉_项目说明书.docx'
SHOT = 'deliverables/shots/s7_report_metro_kpi.png'
LABEL = '图 3-15'


def media_map(path):
    z = zipfile.ZipFile(path)
    rels = z.read('word/_rels/document.xml.rels').decode('utf-8')
    id2t = {i: t.split('/')[-1] for i, t in re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    cam = z.read('word/document.xml').decode('utf-8')
    return id2t, len([n for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/')]), cam.count('<w:drawing')


def main():
    before = [p.text for p in Document(DOCX).paragraphs]
    doc = Document(DOCX)
    target = None
    last_img = None
    for p in doc.paragraphs:
        if 'r:embed' in p._p.xml:
            last_img = p
        if p.text.strip().startswith(LABEL):
            target = last_img
            break
    if target is None:
        print('未找到 %s 前面的图片段落' % LABEL); return
    for r in list(target.runs):
        r._element.getparent().remove(r._element)
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target.add_run().add_picture(SHOT, width=Cm(14.6))
    doc.save(DOCX)
    after = [p.text for p in Document(DOCX).paragraphs]
    print('段落文本一致：', [t for t in before if t.strip()] == [t for t in after if t.strip()])


def report():
    id2t, nmedia, ndraw = media_map(DOCX)
    doc = Document(DOCX)
    pair, last = [], None
    for p in doc.paragraphs:
        m = re.search(r'r:embed="([^"]+)"', p._p.xml)
        if m:
            last = id2t.get(m.group(1), '?')
        if re.match(r'^图 \d+-\d+', p.text.strip()):
            pair.append((p.text.strip().split(chr(12288))[0].replace('图 ', ''), last or '（无图）'))
            last = None
    import collections
    cnt = collections.Counter(m for _, m in pair if m != '（无图）')
    print('媒体 %d ｜ drawing %d ｜ 图注 %d' % (nmedia, ndraw, len(pair)))
    print('重复使用同一张图的图注:', [(a, m) for a, m in pair if cnt[m] > 1] or '无')
    for lab in ('3-13', '3-14', '3-15', '3-16', '3-17'):
        for a, m in pair:
            if a == lab:
                print('   图 %-5s → %s' % (a, m))


if __name__ == '__main__':
    import sys
    report() if len(sys.argv) > 1 and sys.argv[1] == 'report' else main()