# -*- coding: utf-8 -*-
"""把说明书里的界面截图换成分辨率高一倍的版本，只改图片字节，不动正文/排版。

做法：docx 内嵌图与 deliverables/shots 里的 png 原本字节一致；现在 shots 里已是 2x 重拍版，
因此按「尺寸 = 内嵌图 2 倍」做候选、必要时用相关性消歧，逐张替换，再补回缺失的图 3-16。
输出：deliverables/澜脉_项目说明书_高清图.docx（不改原文件）
用法：python src/dsh/2026-09-22_17_sharpen_docx_images.py
"""
import io, os, re, zipfile
import numpy as np
from PIL import Image

SRC = 'deliverables/澜脉_项目说明书.docx'
DST = 'deliverables/澜脉_项目说明书_高清图.docx'
SHOTS = 'deliverables/shots'
MISSING = ('图 3-16', 's9_report_metro_tail.png')


def gray(b):
    return np.asarray(Image.open(io.BytesIO(b)).convert('L'), dtype=np.float32)


def build_new_map():
    m = {}
    for name in sorted(os.listdir(SHOTS)):
        if name.endswith('.png'):
            m[name] = os.path.join(SHOTS, name)
    return m


def pick_media(media_items, new_paths):
    """media_items: [(name, bytes)]；返回 {media_name: shot_name}"""
    sizes = {}
    for name, b in media_items:
        sizes[name] = Image.open(io.BytesIO(b)).size
    assigns, ambiguous = {}, 0
    for name, b in media_items:
        w, h = sizes[name]
        cands = [n for n, p in new_paths.items() if Image.open(p).size == (2 * w, 2 * h)]
        if len(cands) == 1:
            assigns[name] = cands[0]
        elif len(cands) > 1:
            ambiguous += 1
            a = gray(b).ravel()
            a = (a - a.mean()) / (a.std() + 1e-6)
            best, bestc = None, -2
            for c in cands:
                sm = Image.open(new_paths[c]).convert('L').resize((w, h), Image.LANCZOS)
                d = np.asarray(sm, dtype=np.float32).ravel()
                d = (d - d.mean()) / (d.std() + 1e-6)
                cc = float((a * d).mean())
                if cc > bestc:
                    best, bestc = c, cc
            assigns[name] = best
            print('  消歧 %-16s → %-28s 相关 %.3f' % (name, best, bestc))
    print('候选消歧 %d 张' % ambiguous)
    return assigns, sizes


def main():
    new_paths = build_new_map()
    zin = zipfile.ZipFile(SRC)
    media_items = [(n, zin.read(n)) for n in zin.namelist()
                   if n.startswith('word/media/') and n.lower().endswith('.png')]
    assigns, sizes = pick_media(media_items, new_paths)
    zout = zipfile.ZipFile(DST, 'w', zipfile.ZIP_DEFLATED)
    n_rep = 0
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename in assigns:
            shot = assigns[it.filename]
            nb = io.open(new_paths[shot], 'rb').read()
            ns = Image.open(io.BytesIO(nb)).size
            os_ = sizes[it.filename]
            assert ns == (2 * os_[0], 2 * os_[1]), (shot, os_, ns)
            data = nb
            n_rep += 1
            print('  %-16s ← %-28s %sx%s → %sx%s（%.0f → %.0f dpi@14.6cm）' % (
                os.path.basename(it.filename), shot, os_[0], os_[1], ns[0], ns[1],
                os_[0] / (14.6 / 2.54), ns[0] / (14.6 / 2.54)))
        zout.writestr(it, data)
    zout.close(); zin.close()
    print('共替换 %d 张截图' % n_rep)
    fix_missing_figure(new_paths)
    print('输出：%s（%.1f MB）' % (DST, os.path.getsize(DST) / 1048576.0))
    report()


def fix_missing_figure(new_paths):
    from docx import Document
    from docx.shared import Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document(DST)
    for p in doc.paragraphs:
        if p.text.strip().startswith(MISSING[0]):
            np_ = p.insert_paragraph_before()
            np_.alignment = WD_ALIGN_PARAGRAPH.CENTER
            np_.add_run().add_picture(new_paths[MISSING[1]], width=Cm(14.6))
            doc.save(DST)
            print('已补回 %s 配图（%s）' % (MISSING[0], MISSING[1]))
            return True
    print('未找到 %s 图注，未补图' % MISSING[0])
    return False


def report():
    from docx import Document
    z = zipfile.ZipFile(DST)
    cam = z.read('word/document.xml').decode('utf-8')
    docs = Document(DST)
    caps = [p.text for p in docs.paragraphs if re.match(r'^图 \d+-\d+', p.text.strip())]
    last, paired = None, 0
    for p in docs.paragraphs:
        if 'r:embed' in p._p.xml:
            last = True
        if re.match(r'^图 \d+-\d+', p.text.strip()):
            if last:
                paired += 1
            last = None
    print('复核：drawing %d ｜ 图注 %d ｜ 有配图 %d ｜ 表格 %d' % (
        cam.count('<w:drawing'), len(caps), paired, len(docs.tables)))


if __name__ == '__main__':
    main()