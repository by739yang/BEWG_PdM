# -*- coding: utf-8 -*-
"""把配图注入 dsh-ppt 生成的 PPTX（每张图单独成页，插在对应内容页之后）。
用法：python src/dsh/2026-09-22_13_deck_figures.py
"""
import io, os, json, sys
from pptx import Presentation
from pptx.util import Inches, Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
DECK = 'deliverables/ppt/澜脉_商业计划书.pptx'
MANI = 'deliverables/ppt/澜脉_商业计划书.json'

# 目标页标题（子串匹配） -> [(图片, 图注)]
MAP = [
 ('产品：一条链路四段', [('deliverables/figs/figB_pipeline.png', '四段链路：检测 → 诊断 → 剩余寿命 → 维护决策')]),
 ('检测判据长什么样', [('deliverables/figs/figC_criterion.png', '检测判据：双基线分数 + 事件机 + P1/P2/P3 分级')]),
 ('创新一：双基线并行', [('results/2026-09-20/dsh/bsm1_4in1.png', '数字孪生闭环四图：注入退化 → 检测 → RUL → 决策')]),
 ('创新三：设备级比值通道', [('results/2026-09-22/dsh/bsm2_route_b_detect.png', 'BSM2 整厂：原始量检不出，设备级比值通道 9.80 天检出')]),
 ('创新四：机理锚定 RUL', [('results/2026-09-21/dsh/sludge_line2.png', '污泥线闭环：含固率、湿泥饼量、滤液 TSS 与检测分数')]),
 ('创新五：决策层能算钱', [('results/2026-09-21/dsh/sludge_decision.png', '维护策略对比：RUL 驱动相对最优固定周期省 24.5%（占位成本）')]),
 ('验证纪律', [('results/2026-09-18/dsh/figures_4in1.png', '真实数据四图：误报-召回 / 策略成本 / 敏感性 / RUL 散点')]),
 ('产品形态：现场能直接用', [('deliverables/shots/s2_demo_dashboard.png', '十节演示看板（在线可看）：从检测到决策的全部结果与口径')]),
 ('使用方式（二）', [('deliverables/shots/s6_workbench.png', '本地工作台：浏览器上传 CSV → 标定基线 → 出报告，数据不出本机')]),
 ('企业版报告长什么样', [('deliverables/shots/s8_report_metro_alarms.png', '企业版报告·分级告警明细（P1/P2/P3 与处置动作）')]),
 ('现场试点四步走', [('deliverables/figs/figE_poc.png', '现场试点（POC）四步：选设备 → 标定 → 历史回放 → 在线运行')]),
 ('附录：所有数字都能复现', [('deliverables/shots/s4_report_metropt3.png', '样例报告（真实空压机 214 天）：105 个告警，P1 56 / P2 30 / P3 19')]),
]


def slide_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip()
    return ''


def main():
    prs = Presentation(DECK)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    titles = [slide_title(s) for s in prs.slides]
    print('原页数 %d' % len(titles))
    inserted = 0
    for key, figs in MAP:
        idx = None
        for i, t in enumerate(titles):
            if key in t:
                idx = i
                break
        if idx is None:
            print('  [跳过] 找不到标题：%s' % key); continue
        for img, cap in figs:
            if not os.path.exists(img):
                print('  [跳过] 缺图：%s' % img); continue
            s = prs.slides.add_slide(blank)
            # 图：留边距与图注空间
            margin = Emu(int(W * 0.06))
            cap_h = Emu(int(H * 0.10))
            avail_w = W - 2 * margin
            avail_h = H - 2 * margin - cap_h
            from PIL import Image
            iw, ih = Image.open(img).size
            scale = min(avail_w / iw, avail_h / ih)
            w, h = int(iw * scale), int(ih * scale)
            s.shapes.add_picture(img, Emu(int((W - w) / 2)), Emu(int(margin + (avail_h - h) / 2)), width=Emu(w), height=Emu(h))
            tb = s.shapes.add_textbox(margin, Emu(int(H - margin - cap_h + cap_h * 0.25)), avail_w, cap_h)
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run(); run.text = cap
            run.font.size = Pt(14); run.font.color.rgb = RGBColor(0x33, 0x44, 0x55)
            # 移到目标页之后
            sldIdLst = prs.slides._sldIdLst
            ids = list(sldIdLst)
            mv = ids[-1]
            sldIdLst.remove(mv)
            sldIdLst.insert(idx + 1, mv)
            inserted += 1
            titles.insert(idx + 1, key + '（图）')
            idx += 1
    prs.save(DECK)
    prs2 = Presentation(DECK)
    print('插入图页 %d 张 → 新页数 %d' % (inserted, len(prs2.slides)))
    # 校验：每页都有内容
    empty = [i + 1 for i, s in enumerate(prs2.slides) if not any(sh.has_text_frame and sh.text_frame.text.strip() for sh in s.shapes)]
    print('无文字页（应为图片页）：%s' % (empty or '无'))
    print('图片总数：%d' % sum(1 for s in prs2.slides for sh in s.shapes if sh.shape_type == 13))


if __name__ == '__main__':
    main()
