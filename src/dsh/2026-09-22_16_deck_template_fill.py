# -*- coding: utf-8 -*-
"""在官方 PPT 模板上填内容（保留版式与背景）：替换提示文字 + 局部插图。
用法：python src/dsh/2026-09-22_16_deck_template_fill.py
"""
import os, shutil, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
TPL = r'C:\Users\boyi\.dsh\attachments\v1\files\cb\cb8c275ac887da38a5694a0b7ccaabe2a8a246576f6a715dfe4e81a5fac1c0d0\商业计划书-ppt模板.pptx'
OUT = 'deliverables/ppt/澜脉_商业计划书.pptx'
FONT = '微软雅黑'

# 每页（按标题匹配）→ (要写入的正文行, 可选插图)
CONTENT = [
 ('项目概述', [
   '一句话核心：不加任何新传感器，用厂里已有的电流/温度/压力/流量数据，提前发现设备退化并给出维护时机。',
   '项目定位：面向市政污水厂（TO B / TO G）的设备健康管理软件 + 接入标定服务；纯软件、可离线部署、数据不出厂。',
   '对应命题 2-6：交付「检测 → 诊断 → 剩余寿命 → 维护决策」完整闭环，而非单点算法演示。',
   '市场前景：存量污水厂设备普遍缺乏预测性维护手段，数据已在 DCS/SCADA 中，缺方法与工具。'], None),
 ('核心产品/服务', [
   '① 检测：双基线（冻结抓慢漂移 + 短窗自适应抓突变）+ 工况条件化 → 分级告警 P1/P2/P3',
   '② 诊断：CWRU 基准迁移，同记录/跨转速 宏F1 1.000（跨故障尺寸 0.536，边界已标注）',
   '③ 剩余寿命：机理量外推 + 趋势门控 + 拟合窗匹配；孪生场景误差 1.00–4.04 天',
   '④ 决策：RUL 驱动排程 vs 固定周期；C-MAPSS 省约 16%（12/15 格），污泥线脱水机省 24.5%（9/9 格）',
   '交付形态：命令行工具 + 本地网页工作台 + 企业版报告（单文件 HTML）'], 'deliverables/figs/figB_pipeline.png'),
 ('文章及知识产权情况', [
   '短期（1 年内）：完成软件著作权登记（材料已备）；形成 1 篇技术报告/论文初稿；整理公开数据集上的复现脚本。',
   '中期（3 年）：申请发明专利（双基线+工况条件化的检测方法、设备级物料平衡比值监测方法）；参与行业标准/团体标准讨论。',
   '长期（5 年）：形成「设备健康管理」方法学与工具链，向供水厂、排水管网泵站横向扩展。',
   '当前状态：代码与文档均可现场演示；所有数字可复现（一键复现 39 阶段 + 数字总检 30/30）。'], None),
 ('产品迭代计划', [
   '短期（1 年内）：接入首个污水厂试点（1 台设备）→ 用真实数据校准阈值与工况分层；补齐诊断到水务设备语义的映射。',
   '中期（3 年）：扩展至鼓风机/脱水机/泵等多设备类型；增加多源冗余（在线密度计 + 实验室）与独立漂移监控模块。',
   '长期（5 年）：从「单厂工具」演进为「集团级设备健康平台」，支持多厂复制与统一看板。'], None),
 ('项目亮点', [
   '创新性：三层防误报（双基线 / 工况条件化 / 同期同工况对拍）+ 设备级物料平衡比值监测 + 机理锚定 RUL + 诚实工程（四轮独立复核、数字总检）。',
   '社会价值：减少非计划停机与溢流风险；泥饼含固率不达标的合规预警；不新增传感器，不增加现场改造负担。',
   '商业价值：按设备数订阅，轻资产、不垫资；可复制到供水与管网泵站。',
   '关键证据：整厂尺度原始量 86.11 天仍检不出 → 设备级比值通道 9.80 天检出、提前 37.47 天。'], 'results/2026-09-22/dsh/bsm2_route_b_detect.png'),
 ('行业背景', [
   '存量巨大：城市及县城污水处理厂约 4816 座（2023 城乡建设统计年鉴，待原文核实）。',
   '政策导向：双碳与污水资源化持续加压，运维成本与合规要求同步提高。',
   '现状：关键设备以「定期大修 + 坏了才修」为主，缺乏预测性维护手段。',
   '数据条件：运行数据已在 DCS/SCADA 里，缺的是方法与工具（对齐本次命题）。'], None),
 ('目标市场', [
   '使用者：厂级设备科 / 运维班长（依据 P1/P2/P3 决定动作）。',
   '付费方：水务集团 / 地方水司的生产技术部或设备管理部门。',
   '决策链：厂级试点 → 集团技术部评估 → 集团统一部署。',
   '典型客户特征：有明确运维考核指标（非计划停机、能耗、药耗），且愿意做小范围试点。'], None),
 ('细分市场规模', [
   '口径一：存量厂数 × 单厂关键设备数（鼓风机 / 脱水机 / 提升泵 / 回流泵）。',
   '口径二：单厂年度运维预算中「非计划停机损失 + 备件与人工」占比。',
   '口径三：对标工业预测性维护软件单价区间 × 渗透率。',
   '说明：三种口径均需公开原文核实后再填报，未核实前不外报具体金额。'], None),
 ('直接竞争对手与替代品分析', [
   '国际预测性维护平台：功能全、算法成熟，但价格高、需联网/云，数据出厂的合规顾虑大。',
   '通用工业 IoT 平台：擅长工艺全景与看板，但看不见设备级退化（本作品实测：沼气/能耗/出水对脱水机退化几乎无响应，±0.03%）。',
   '传统咨询式检修：经验丰富但不可复制、不可量化、无法持续在线。',
   '替代品：定期大修制度本身 —— 我们不以「取代它」为卖点，而是让检修计划按设备实际状态排。'], None),
 ('项目竞争优势', [
   '机制层面把误报压下去：同召回下误报事件比对照实现少 2.8 倍（三层防误报）。',
   '能发现别人看不见的设备级退化：整厂原始量检不出 → 设备级比值量 9.80 天检出。',
   '现场能自己接数据：命令行 + 本地工作台 + 企业版报告，零外部依赖、可离线、数据不出厂。',
   '数字可核查：四轮跨实现独立复核 + 冻结前数字总检 30/30 + 一键复现 39 阶段。'], None),
 ('市场进入策略', [
   '第一步：免费试点 1 台设备（建议鼓风机或污泥脱水机），出健康基线报告。',
   '第二步：用含故障的历史段回放，给出「提前量 / 误报数」对照报告 —— 用数据而不是演示说服客户。',
   '第三步：单厂多设备推广（鼓风机 / 脱水机 / 泵），再进入集团多厂复制。',
   '切入渠道：水务集团技术部、市政设计院、设备厂商配套、大赛产业对接会。'], 'deliverables/figs/figE_poc.png'),
 ('营销渠道', [
   '集团渠道：通过水务集团生产技术部做集团级试点与统一采购。',
   '设计院渠道：在新建/改造项目中作为数字化配套方案进入设计清单。',
   '设备厂商渠道：与鼓风机、脱水机厂商配套销售（我们提供软件与标定服务）。',
   '大赛渠道：命题方北控水务的产业对接会与行业创新成果展（本次赛事提供的资源）。'], None),
 ('盈利模式', [
   '收入一：按设备数收取软件订阅年费（含更新与技术支持）。',
   '收入二：一次性接入标定服务（数据接入 + 基线标定 + 首份健康报告）。',
   '收入三（可选）：集团级部署、培训与定制开发。',
   '定价依据：以「单台设备一次非计划停机的损失」为价值上限，按设备数与服务深度分档；具体价格待企业成本参数标定。'], None),
 ('财务预测', [
   '成本结构：研发与维护（人力为主）、试点交付（差旅与实施）、部署与支持（轻资产、无硬件）。',
   '收入结构：试点转付费率 × 单厂设备数 × 年费（+ 一次性接入标定服务）。',
   '关键指标口径：毛利率、客户获取成本、续费率、单厂回本周期。',
   '说明：本项目所有金额均为占位值，需用企业现场成本参数（失效损失 / 计划更换费用 / 剩余寿命折算）标定后再填报。'], None),
 ('融资计划', [
   '当前阶段：无外部融资需求 —— 纯软件研发，自筹 + 大赛奖金即可覆盖。',
   '若启动试点：资金用于试点实施与交付人力（非重资产投入）。',
   '融资用途（如发生）：产品化（多设备类型、多源冗余与漂移监控）、试点交付、市场与渠道建设。'], None),
 ('当前进展', [
   '技术侧已冻结待命：四段链路 + 两套机理仿真（BSM1 / BSM2 整厂）+ 两套真实数据验证（水泵台架 / 地铁空压机 214 天）。',
   '可信度建设：四轮跨实现独立复核（问题全部回写）+ 冻结前数字总检 30/30 + 一键复现 39 阶段。',
   '交付物：项目说明书（Word）、商业计划书（PPT）、本地工作台、企业版报告、线上演示页。',
   '待补：真实污水厂试点数据（正在申请）、企业成本参数。'], 'deliverables/shots/s6_workbench.png'),
 ('未来里程碑', [
   '10/07：技术冻结（可复现包 + 文档 + 冻结 tag）。',
   '11 月：大赛产业对接会，争取 1 个污水厂试点意向。',
   '2027 Q1：首个污水厂试点上线（1 台设备、3–6 个月对照）。',
   '2027 年内：扩展至单厂多设备，并完成首批付费客户签约。'], None),
 ('主要风险', [
   '技术风险：真实厂数据缺失导致阈值无法现场标定；传感器标定漂移可能造成假 RUL。',
   '市场风险：客户对「预测性维护」付费意愿需验证；集团采购周期长。',
   '数据风险：现场数据质量（缺失、工况列不规范）影响检出效果。',
   '团队风险：缺少给排水/环境专业成员与行业顾问，指导老师待定。'], None),
 ('应对措施', [
   '技术：试点前用历史段回放验证；标定漂移用多源冗余（在线密度计 + 实验室）+ 独立漂移监控压制；RUL 只挂机理量并带趋势门控。',
   '市场：先做免费试点，用「提前量 / 误报数」对照报告说话；以单厂为最小付费单元降低决策门槛。',
   '数据：接入前先体检（缺失率、零方差、参考窗趋势），并要求多段健康期标定。',
   '团队：优先补齐 1 名给排水/环境专业成员与 1 名市场成员，并争取高校实验室与产业导师支持。'], None),
 ('团队构成', [
   '项目负责人：<姓名>（算法与工程实现、材料撰写）。',
   '拟补成员 1：给排水 / 环境工程专业（把关专业说法与现场流程）。',
   '拟补成员 2：市场与商务（市场调研、客户接触、PPT 与路演）。',
   '指导老师：<待定>；合作资源：高校实验室 + 大赛产业导师 + 命题方产业对接渠道。'], None),
 ('附录', [
   '交付物：项目说明书（Word，38 图）、商业计划书（本 PPT）、本地工作台、企业版报告。',
   '在线演示：落地页 / 十节看板 / 交互式数字孪生演示台 / 两份真实样例报告（GitHub Pages）。',
   '可复现：run_all（快速 9 阶段 / 全量 39 阶段）+ 四轮独立复核台账 + 冻结前数字总检 30/30。',
   '口径纪律：仿真一律标注「仿真」；成本参数为占位值；未标定前不对外给金额。'], 'deliverables/shots/s8_report_metro_alarms.png'),
]

def set_font(run, size=14, bold=False, color='1B2733', name=FONT):
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set('typeface', name)

def texts_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            yield sh, sh.text_frame.text.strip()

def fill_lines(tf, lines, size=13.5):
    tf.clear()
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run(); run.text = ('· ' + ln) if not ln.startswith(('①', '②', '③', '④', '·')) else ln
        set_font(run, size)
        p.space_after = Pt(6)

def main():
    prs = Presentation(TPL)
    slides = list(prs.slides)
    # 封面
    for sh in slides[0].shapes:
        if sh.has_text_frame and '项目名称' in sh.text_frame.text:
            tf = sh.text_frame; tf.clear()
            r = tf.paragraphs[0].add_run(); r.text = '澜脉 · 污水厂设备 AI 预测性维护系统'
            set_font(r, 28, True, '12507B')
            p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = '商业计划书（创意转化组 · 命题方向 2-6）'
            set_font(r2, 16, False, '31465A')
    filled, imgs = [], 0
    for key, lines, img in CONTENT:
        for s in slides:
            ttl = [t for _, t in texts_of(s)]
            if not ttl or key not in ttl[0]:
                continue
            # 找到「提示文字」形状：AUTO_SHAPE，或第 2 个文本框
            hint = None
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip() and sh.text_frame.text.strip() != ttl[0]:
                    if str(sh.shape_type).startswith('AUTO_SHAPE') or hint is None:
                        hint = sh
            if hint is None:
                tb = s.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(6.0), Inches(4.2))
                hint = tb
            hint.width = Inches(5.9); hint.height = Inches(4.3); hint.top = Inches(1.35); hint.left = Inches(0.75)
            fill_lines(hint.text_frame, lines, 13.5 if len(lines) <= 4 else 12.5)
            hint.text_frame.word_wrap = True
            if img and os.path.exists(img):
                iw, ih = Image.open(img).size
                box_w, box_h = Inches(6.1), Inches(4.2)
                sc = min(box_w / iw, box_h / ih)
                w, h = int(iw * sc), int(ih * sc)
                s.shapes.add_picture(img, Inches(6.9) + Emu(int((box_w - w) / 2)), Inches(1.5) + Emu(int((box_h - h) / 2)), width=Emu(w), height=Emu(h))
                imgs += 1
            filled.append(key)
            break
    # 删掉「请注意」使用提示页（模板自带的填写说明）
    last = list(prs.slides)
    if last and '请注意' in ' '.join(t for _, t in texts_of(last[-1])):
        lst = prs.slides._sldIdLst; ids = list(lst); lst.remove(ids[-1])
        print('已删除模板自带的「请注意」说明页')
    os.makedirs('deliverables/ppt', exist_ok=True)
    prs.save(OUT)
    chk = Presentation(OUT)
    print('模板填充完成：写入 %d 页内容 ｜ 插图 %d 张 ｜ 最终页数 %d' % (len(filled), imgs, len(chk.slides)))
    miss = [k for k, _, _ in CONTENT if k not in filled]
    print('未匹配的板块：%s' % (miss or '无'))
    print('产物：%s（%.0f KB）' % (OUT, os.path.getsize(OUT) / 1024))

if __name__ == '__main__':
    main()
