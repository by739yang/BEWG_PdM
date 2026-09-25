# -*- coding: utf-8 -*-
"""PPT 定稿：从临时目录取 pptx → 注入配图 → 去重复收尾页 → 校验页数 ≤30 → 只保留 pptx。
用法：python src/dsh/2026-09-22_14_deck_finalize.py
"""
import os, shutil, glob
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
TMP = 'deliverables/_tmp_deck'
DST_DIR = 'deliverables/ppt'
DST = os.path.join(DST_DIR, '澜脉_商业计划书.pptx')
MAXPAGES = 30

# 目标页标题（子串） -> (图片, 图注)
MAP = [
 ('检测判据长什么样', 'deliverables/figs/figC_criterion.png', '检测判据：双基线分数 + 事件机 + P1/P2/P3 分级'),
 ('创新三：设备级比值通道', 'results/2026-09-22/dsh/bsm2_route_b_detect.png', 'BSM2 整厂：原始量检不出，设备级比值通道 9.80 天检出'),
 ('使用方式：四条命令', 'deliverables/shots/s6_workbench.png', '本地工作台：浏览器上传 CSV → 标定基线 → 出报告（数据不出本机）'),
 ('市场进入策略', 'deliverables/figs/figE_poc.png', '现场试点（POC）四步：选设备 → 标定 → 历史回放 → 在线运行'),
]


def title_of(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip()
    return ''


def main():
    src = glob.glob(os.path.join(TMP, '*.pptx'))
    if not src:
        raise SystemExit('临时目录没有 pptx：%s' % TMP)
    os.makedirs(DST_DIR, exist_ok=True)
    shutil.copy(src[0], DST)
    prs = Presentation(DST)
    blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    W, H = prs.slide_width, prs.slide_height
    titles = [title_of(s) for s in prs.slides]
    print('注入前页数 %d' % len(titles))
    added = 0
    for key, img, cap in MAP:
        idx = next((i for i, t in enumerate(titles) if key in t), None)
        if idx is None:
            print('  [跳过] 未找到：%s' % key); continue
        if not os.path.exists(img):
            print('  [跳过] 缺图：%s' % img); continue
        s = prs.slides.add_slide(blank)
        margin = Emu(int(W * 0.055)); cap_h = Emu(int(H * 0.095))
        avail_w, avail_h = W - 2 * margin, H - 2 * margin - cap_h
        iw, ih = Image.open(img).size
        sc = min(avail_w / iw, avail_h / ih)
        w, h = int(iw * sc), int(ih * sc)
        s.shapes.add_picture(img, Emu(int((W - w) / 2)), Emu(int(margin + (avail_h - h) / 2)), width=Emu(w), height=Emu(h))
        tb = s.shapes.add_textbox(margin, Emu(int(H - margin - cap_h + cap_h * 0.3)), avail_w, cap_h)
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = cap
        run.font.size = Pt(13); run.font.color.rgb = RGBColor(0x33, 0x44, 0x55)
        lst = prs.slides._sldIdLst; ids = list(lst); mv = ids[-1]
        lst.remove(mv); lst.insert(idx + 1, mv)
        titles.insert(idx + 1, key + '（图）'); idx += 1; added += 1
    # 去重复收尾页（dsh-ppt 自带「谢谢」）
    t = [title_of(s) for s in prs.slides]
    if len(t) >= 2 and t[-1].strip() in ('谢谢', '谢谢观看', 'THANKS', 'Thanks'):
        lst = prs.slides._sldIdLst; ids = list(lst); lst.remove(ids[-1])
        print('已删除重复收尾页')
    prs.save(DST)
    chk = Presentation(DST)
    n = len(chk.slides)
    print('注入图页 %d 张 → 最终页数 %d（上限 %d）%s' % (added, n, MAXPAGES, 'OK' if n <= MAXPAGES else 'OVER'))
    print('图片总数 %d' % sum(1 for s in chk.slides for sh in s.shapes if sh.shape_type == 13))
    shutil.rmtree(TMP, ignore_errors=True)
    print('临时目录已清理：', TMP)
    print('产物：', DST)
    print('目录内容：', sorted(os.listdir(DST_DIR)))


if __name__ == '__main__':
    main()
